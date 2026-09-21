#!/usr/bin/env python3
"""Run the synthetic QMS screen through the local adapter contract.

This is a development fixture, not a production connector. It exercises the
same boundaries a local ERP/QMS adapter must use: read a dedicated browser
page, normalize records, render exactly one DOCX, and submit a hashed manifest
to the draft-first harness.
"""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import sys
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

import harness  # noqa: E402


try:
    from docx import Document
    from playwright.sync_api import sync_playwright
except ImportError as exc:  # pragma: no cover - exercised by the setup error path
    Document = None
    sync_playwright = None
    IMPORT_ERROR = exc
else:
    IMPORT_ERROR = None


EXPECTED_COLUMNS = ("record_id", "title", "status", "owner", "updated_at")


def require_dependencies() -> None:
    if IMPORT_ERROR is not None:
        raise RuntimeError(
            "fixture dependencies are missing; install with "
            "python3 -m pip install -r agent-harness/requirements-fixture.txt "
            "and then run `python3 -m playwright install chromium`"
        ) from IMPORT_ERROR


def extract_qms_screen(html_path: Path) -> tuple[str, str, list[dict[str, Any]], list[dict[str, Any]]]:
    """Read the fixture through a real headless Chromium page."""
    require_dependencies()
    page_url = html_path.resolve().as_uri()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            page = browser.new_page()
            page.goto(page_url, wait_until="load")
            root = page.locator("main[data-source-system]").first
            source_system = root.get_attribute("data-source-system") or ""
            captured_at = root.get_attribute("data-captured-at") or ""
            headers = [text.strip() for text in page.locator("#qms-open-issues thead th").all_text_contents()]
            if tuple(headers) != EXPECTED_COLUMNS:
                raise ValueError(f"unexpected QMS columns: {headers}")
            rows = page.locator("#qms-open-issues tbody tr")
            records: list[dict[str, Any]] = []
            seen: set[str] = set()
            for index in range(rows.count()):
                row = rows.nth(index)
                record_id = (row.get_attribute("data-record-id") or "").strip()
                values = [text.strip() for text in row.locator("td").all_text_contents()]
                if len(values) != len(EXPECTED_COLUMNS):
                    raise ValueError(f"row {index + 1} has {len(values)} cells")
                record = dict(zip(EXPECTED_COLUMNS, values))
                if record["record_id"] != record_id or not record_id:
                    raise ValueError(f"row {index + 1} has inconsistent record_id")
                if record_id in seen:
                    raise ValueError(f"duplicate record_id: {record_id}")
                if any(not value for value in record.values()):
                    raise ValueError(f"blank value in record: {record_id}")
                seen.add(record_id)
                record["source_ref"] = f"qms://issue/{record_id}"
                record["evidence_selector"] = f"#qms-open-issues tr[data-record-id='{record_id}']"
                records.append(record)
        finally:
            browser.close()
    checks = [
        {"name": "source_system", "passed": source_system == "QMS"},
        {"name": "required_columns", "passed": True},
        {"name": "non_empty_records", "passed": bool(records)},
        {"name": "duplicate_ids", "passed": len({record["record_id"] for record in records}) == len(records)},
    ]
    try:
        datetime.fromisoformat(captured_at.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"invalid captured_at: {captured_at}") from exc
    return source_system, captured_at, records, checks


def render_docx(
    output_path: Path,
    source_system: str,
    captured_at: str,
    records: list[dict[str, Any]],
) -> None:
    require_dependencies()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("QMS Daily Draft", level=1)
    document.add_paragraph(f"Source: {source_system} | Captured: {captured_at}")
    document.add_paragraph("Draft only. No ERP/QMS write or external send was performed.")
    table = document.add_table(rows=1, cols=len(EXPECTED_COLUMNS))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, EXPECTED_COLUMNS):
        cell.text = heading
    for record in records:
        cells = table.add_row().cells
        for cell, heading in zip(cells, EXPECTED_COLUMNS):
            cell.text = record[heading]
    document.add_heading("Evidence", level=2)
    for record in records:
        document.add_paragraph(f"{record['record_id']}: {record['source_ref']}")
    document.save(output_path)


def reopen_docx(output_path: Path, expected_records: int) -> dict[str, Any]:
    require_dependencies()
    document = Document(output_path)
    table_rows = len(document.tables[0].rows) if document.tables else 0
    passed = bool(document.paragraphs) and table_rows == expected_records + 1
    return {"name": "docx_reopen_and_row_count", "passed": passed}


def sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_output(call: list[str]) -> dict[str, Any]:
    buffer = io.StringIO()
    with contextlib.redirect_stdout(buffer):
        code = harness.main(call)
    lines = [line for line in buffer.getvalue().splitlines() if line.strip()]
    if code != 0 or not lines:
        raise RuntimeError(buffer.getvalue() or f"harness command failed: {call}")
    return json.loads(lines[-1])


def run_demo(config_path: Path, job_id: str, html_path: Path, claim: bool = False) -> dict[str, Any]:
    config = harness.load_config(config_path)
    if claim:
        claim_result = _json_output(["--config", str(config_path), "run-once", "--claim"])
        if claim_result.get("job_id") != job_id:
            raise RuntimeError(f"claimed {claim_result.get('job_id')} instead of {job_id}")
    with harness.connect(config) as connection:
        row = connection.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row is None:
            raise ValueError(f"unknown job: {job_id}")
        job = harness.row_json(row)
    if job["status"] != "running":
        raise ValueError(f"job must be running; use --claim for a queued job (current: {job['status']})")
    source_system, captured_at, records, checks = extract_qms_screen(html_path)
    artifact_path = harness.path_from_config(config, "artifact_dir") / f"{job_id}.docx"
    render_docx(artifact_path, source_system, captured_at, records)
    checks.append(reopen_docx(artifact_path, len(records)))
    manifest_path = harness.path_from_config(config, "state_dir") / f"{job_id}.manifest.json"
    manifest = {
        "schema_version": harness.MANIFEST_VERSION,
        "job_id": job_id,
        "source_system": source_system,
        "captured_at": captured_at,
        "records": records,
        "checks": checks,
        "artifact": {
            "path": str(artifact_path),
            "format": job["output_format"],
            "sha256": sha256_file(artifact_path),
        },
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    result = _json_output(
        [
            "--config",
            str(config_path),
            "record-result",
            "--job-id",
            job_id,
            "--manifest",
            str(manifest_path),
        ]
    )
    return {"job_id": job_id, "records": len(records), "artifact": str(artifact_path), "manifest": str(manifest_path), "result": result}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Synthetic QMS browser-to-DOCX adapter")
    parser.add_argument("--config", required=True)
    parser.add_argument("--job-id", required=True)
    parser.add_argument("--html", default=str(ROOT / "agent-harness" / "fixtures" / "qms_daily.html"))
    parser.add_argument("--claim", action="store_true", help="claim the next queued job before extraction")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        print(json.dumps(run_demo(Path(args.config), args.job_id, Path(args.html), args.claim), ensure_ascii=False, indent=2))
        return 0
    except (OSError, RuntimeError, ValueError, json.JSONDecodeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
